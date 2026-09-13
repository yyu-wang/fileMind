"""doc_extract 单元测试：四种二进制格式抽取、损坏文件、上限截断、类型判定。

PDF 测试样本由 :func:`_build_pdf` 手写构造（正确 xref 偏移的最小单页文档），
不引入 reportlab 等仅测试用的重型依赖。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.services import doc_extract


def _build_pdf(content: str) -> bytes:
    """构造包含单个文本对象的合法单页 PDF（xref 偏移精确计算）。

    Args:
        content: 页面上渲染的 ASCII 文本。

    Returns:
        pypdf 可直接解析的 PDF 字节。
    """
    buf = bytearray()
    offsets: dict[int, int] = {}

    def _write(num: int, body: bytes) -> None:
        offsets[num] = len(buf)
        buf.extend(f"{num} 0 obj\n".encode("ascii"))
        buf.extend(body)
        buf.extend(b"\nendobj\n")

    _write(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    _write(2, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    _write(
        3,
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
    )
    escaped = content.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 14 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    stream_obj = f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream"
    _write(4, stream_obj)
    _write(5, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    xref_pos = len(buf)
    buf.extend(b"xref\n0 6\n")
    buf.extend(b"0000000000 65535 f \n")
    for i in range(1, 6):
        buf.extend(f"{offsets[i]:010d} 00000 n \n".encode("ascii"))
    buf.extend(
        (f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n").encode("ascii")
    )
    return bytes(buf)


def test_is_binary_document() -> None:
    assert doc_extract.is_binary_document(Path("a.PDF"))
    assert doc_extract.is_binary_document(Path("/x/report.docx"))
    assert doc_extract.is_binary_document(Path("b.xlsx"))
    assert doc_extract.is_binary_document(Path("c.pptx"))
    assert not doc_extract.is_binary_document(Path("a.txt"))
    assert not doc_extract.is_binary_document(Path("a.png"))
    assert not doc_extract.is_binary_document(Path("noext"))


def test_extract_pdf_returns_text(tmp_path: Path) -> None:
    p = tmp_path / "sample.pdf"
    p.write_bytes(_build_pdf("Hello FileMind PDF"))
    text = doc_extract.extract_document_text(p)
    assert "Hello FileMind PDF" in text


def test_extract_docx_includes_table(tmp_path: Path) -> None:
    p = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("季度报告正文")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "数值"
    doc.save(str(p))

    text = doc_extract.extract_document_text(p)
    assert "季度报告正文" in text
    assert "指标 | 数值" in text


def test_extract_xlsx_includes_cells(tmp_path: Path) -> None:
    p = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["季度", "2026Q1"])
    ws.append(["收入", 123])
    wb.save(str(p))

    text = doc_extract.extract_document_text(p)
    assert "2026Q1" in text
    assert "123" in text


def test_extract_pptx_includes_text_frame(tmp_path: Path) -> None:
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text = "幻灯片演示正文"
    prs.save(str(p))

    text = doc_extract.extract_document_text(p)
    assert "幻灯片演示正文" in text


@pytest.mark.parametrize("name", ["bad.pdf", "bad.docx", "bad.xlsx", "bad.pptx"])
def test_corrupt_document_raises(tmp_path: Path, name: str) -> None:
    p = tmp_path / name
    p.write_bytes(b"garbage-bytes-not-a-real-document")
    with pytest.raises(doc_extract.DocumentExtractError):
        doc_extract.extract_document_text(p)


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    p = tmp_path / "note.txt"
    p.write_text("plain text", encoding="utf-8")
    with pytest.raises(doc_extract.DocumentExtractError):
        doc_extract.extract_document_text(p)


def test_extract_truncates_oversize_docx(tmp_path: Path) -> None:
    p = tmp_path / "big.docx"
    doc = Document()
    doc.add_paragraph("a" * (doc_extract.MAX_EXTRACT_CHARS + 500))
    doc.save(str(p))

    text = doc_extract.extract_document_text(p)
    assert len(text) <= doc_extract.MAX_EXTRACT_CHARS
